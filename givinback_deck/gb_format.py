# -*- coding: utf-8 -*-
"""
Givin' Back 資料フォーマット・ツールキット

`ジンジブ様_CoreUniqオンボーディング・POC資料` の修正版（田中さん手入れ版）から
採寸したスタイルを、そのまま他の資料にも使える形にまとめたもの。

使い方:

    from gb_format import Deck, NAVY, CORAL

    deck = Deck()
    deck.cover("見出し", "サブ見出し", eyebrow="◯◯様｜案件名")

    s = deck.slide("DATA｜見出し", "リード文")
    s.card(ML, 1.86, 6.0, 3.9, accent=NAVY)
    s.closing("締めの一文")

    last = deck.gradient_slide("NEXT ACTION｜見出し", "リード文")
    deck.save("out.pptx")

寸法はすべて inch。スライドは 16:9（13.333 × 7.5 inch = 960 × 540 pt）。
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))

# ------------------------------------------------------------------ カラー
NAVY = RGBColor(0x14, 0x4B, 0x7F)   # 基調
NAVY_DP = RGBColor(0x0E, 0x35, 0x5C)   # 表紙グラデーションの濃い側
INK = RGBColor(0x1E, 0x29, 0x3B)   # 本文
GRAY = RGBColor(0x6B, 0x77, 0x87)   # 補足
GRAY_L = RGBColor(0x93, 0x9E, 0xAC)   # フッター・矢印
CORAL = RGBColor(0xF9, 0x61, 0x67)   # アクセント
CORAL_D = RGBColor(0xE0, 0x45, 0x4B)   # アクセント（文字用の濃い側）
PINK_L = RGBColor(0xFD, 0xEC, 0xED)   # 淡ピンク面
BLUE_L = RGBColor(0xEE, 0xF3, 0xF8)   # 淡ブルー面
TRACK = RGBColor(0xEE, 0xF1, 0xF5)   # バーの下地
BLUE_M = RGBColor(0x6F, 0x9A, 0xC8)   # サブのバー
CHEVRON = RGBColor(0xE8, 0xE8, 0xE8)   # 大矢印
LINE = RGBColor(0xD8, 0xE0, 0xEA)   # 罫線・カード枠
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xEF, 0xF4, 0xFA)   # 表紙の従属テキスト
ROSE_L = RGBColor(0xFF, 0xD9, 0xDB)   # 表紙のアイブロウ

# ------------------------------------------------------------------ フォント
# 修正版で使われているのは Zen Kaku Gothic Antique（Google Fonts）。
# PowerPoint 環境で未インストールの場合は FONT を "Meiryo" などに変えて再生成する。
FONT = "Zen Kaku Gothic Antique"

# ------------------------------------------------------------------ グリッド
SW, SH = 13.333, 7.5
ML, MR = 0.62, 12.71          # 左右マージン
CW = MR - ML                  # コンテンツ幅 12.09

HDR_Y = 0.40                  # ヘッダー見出しの上端
RULE_Y = 0.95                 # ヘッダー罫線
LEAD_X = 0.82                 # リード文の左端（マージンより少し内側）
LEAD_Y = 1.17
BODY_TOP = 1.86               # 本文エリアの上端
CLOSE_X = 0.88                # 締め文の左端
CLOSE_Y = 6.15
FOOT_RULE_Y = 7.02
PAD = 0.26                    # カード内の余白

# ------------------------------------------------------------------ 級数
SZ_TITLE = 21                 # ヘッダー見出し
SZ_LEAD = 24                  # リード文
SZ_CLOSE = 24                 # 締め文
SZ_CARD_T = 20                # カード見出し
SZ_BODY = 14                  # カード本文
SZ_ROW = 19.5                 # 一覧行
SZ_NUM = 24                   # 大きい連番
SZ_FOOT = 7.5

FOOTER = "© Givin' Back Inc. All Rights Reserved.  |  Confidential"

LOGO = os.path.join(HERE, "lock.png")        # 横組みロゴ（濃色）
LOGO_W = os.path.join(HERE, "lockw.png")     # 横組みロゴ（白）


# ==================================================================== Slide
class Slide(object):
    """1枚のスライド。プリミティブを生やして組み立てる。"""

    def __init__(self, deck, sld, dark=False):
        self.deck = deck
        self.sld = sld
        self.dark = dark          # グラデーション面かどうか

    # ------------------------------------------------------------ 基本部品
    def text(self, x, y, w, h, lines, size=SZ_BODY, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.25):
        """lines は str か [(文字列, {size/bold/color/align/spacing}), ...]"""
        box = self.sld.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        if isinstance(lines, str):
            lines = [lines]
        for i, item in enumerate(lines):
            content, opt = item if isinstance(item, tuple) else (item, {})
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = opt.get("align", align)
            p.line_spacing = opt.get("spacing", spacing)
            run = p.add_run()
            run.text = content
            f = run.font
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
            f.name = FONT
            rPr = run._r.get_or_add_rPr()
            for tag in ("ea", "cs"):
                rPr.append(rPr.makeelement(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag,
                    {"typeface": FONT}))
        return box

    def shape(self, x, y, w, h, kind=MSO_SHAPE.RECTANGLE, fill=None,
              line=None, line_w=1.0, adj=None):
        sp = self.sld.shapes.add_shape(kind, Inches(x), Inches(y),
                                       Inches(w), Inches(h))
        if adj is not None:
            try:
                sp.adjustments[0] = adj
            except (IndexError, ValueError):
                pass
        if fill is None:
            sp.fill.background()
        else:
            sp.fill.solid()
            sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line
            sp.line.width = Pt(line_w)
        sp.shadow.inherit = False
        sp.text_frame.word_wrap = True
        return sp

    def rule(self, x, y, w, color=LINE, h=0.014):
        return self.shape(x, y, w, h, fill=color)

    def picture(self, path, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return self.sld.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    # ------------------------------------------------------------ 構造要素
    def lead(self, body, color=None, size=SZ_LEAD):
        """ヘッダー直下の大見出し（罫線バーは付けない）"""
        return self.text(LEAD_X, LEAD_Y, CW - 0.20, 0.50, body,
                         size=size, bold=True,
                         color=color or (WHITE if self.dark else NAVY))

    def closing(self, body, accent=False, y=CLOSE_Y, size=SZ_CLOSE):
        """最下部の締めの一文。accent=True でコーラル。"""
        return self.text(CLOSE_X, y, CW - 0.26, 0.60, body, size=size, bold=True,
                         color=CORAL_D if accent else NAVY)

    def card(self, x, y, w, h, accent=None, side=None, tint=WHITE,
             border=LINE, bar=0.10):
        """白カード。accent=色 で上辺に帯、side=色 で左辺に帯。"""
        self.shape(x, y, w, h, MSO_SHAPE.ROUNDED_RECTANGLE, fill=tint,
                   line=border, line_w=1.0, adj=0.035)
        if accent:
            self.shape(x, y, w, bar, fill=accent)
        if side:
            self.shape(x, y, bar, h, fill=side)

    def pill(self, x, y, w, label, fill=NAVY, size=17, h=0.40,
             color=WHITE, align=PP_ALIGN.LEFT, indent=PAD):
        """見出し帯。align を CENTER にすると中央寄せの大きい帯にも使える。"""
        self.shape(x, y, w, h, fill=fill)
        tx = x + (0 if align == PP_ALIGN.CENTER else indent)
        tw = w if align == PP_ALIGN.CENTER else w - indent * 2
        self.text(tx, y + (h - size / 72.0 * 1.35) / 2, tw, h,
                  label, size=size, bold=True, color=color, align=align)

    def num_badge(self, x, y, d, n, fill=NAVY, size=SZ_NUM):
        self.shape(x, y, d, d, MSO_SHAPE.OVAL, fill=fill)
        self.text(x, y + (d - size / 72.0 * 1.4) / 2, d, d, str(n),
                  size=size, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    def checks(self, x, y, w, items, pitch=0.68, size=SZ_BODY, inline=False):
        """チェック項目。inline=True で『✓文言』をコーラル太字の1行にする。"""
        for i, it in enumerate(items):
            yy = y + i * pitch
            if inline:
                self.text(x, yy, w, 0.34, "✓" + it, size=size, bold=True,
                          color=CORAL)
            else:
                self.text(x, yy, 0.26, 0.30, "✓", size=11.5, bold=True,
                          color=CORAL)
                self.text(x + 0.26, yy - 0.03, w - 0.26, 0.32, it, size=size,
                          color=INK)

    def rows(self, x, y, w, items, pitch=0.58, h=0.46, size=SZ_ROW,
             numbered=True, fill=BLUE_L):
        """淡ブルーの一覧行。numbered=True で先頭にネイビーの小さい連番。"""
        for i, it in enumerate(items):
            yy = y + i * pitch
            self.shape(x, yy, w, h, fill=fill)
            tx = x + PAD
            if numbered:
                d = 0.33
                self.num_badge(x + 0.13, yy + (h - d) / 2, d, i + 1, size=13)
                tx = x + 0.64
            self.text(tx, yy + (h - size / 72.0 * 1.35) / 2, w - (tx - x) - PAD,
                      h, it, size=size, color=INK)

    def agenda(self, items, y=BODY_TOP, header="アジェンダ", indent=1.42,
               head_h=0.40, row_h=0.60, size=18):
        """アジェンダ一覧（ネイビーの見出し行＋交互の淡ブルー行）"""
        self.shape(ML, y, CW, head_h, fill=NAVY)
        self.text(ML + indent, y + (head_h - size / 72.0 * 1.35) / 2, CW - indent,
                  head_h, header, size=size, bold=True, color=WHITE)
        for i, it in enumerate(items):
            yy = y + head_h + i * row_h
            if i % 2 == 0:
                self.shape(ML, yy, CW, row_h, fill=BLUE_L)
            self.rule(ML, yy + row_h - 0.008, CW, h=0.008)
            self.text(ML + indent, yy + (row_h - size / 72.0 * 1.35) / 2,
                      CW - indent - PAD, row_h, it, size=size, bold=True,
                      color=NAVY)
        return y + head_h + len(items) * row_h

    def bars(self, x, y, w, items, label_w=2.55, pitch=0.60, bar_h=0.24,
             vmax=None, highlight=2, unit="%"):
        """横棒グラフ。items は [(ラベル, 値), ...]。上位 highlight 本をコーラルに。"""
        vmax = vmax or max(v for _, v in items) * 1.09
        bx = x + label_w
        bw = w - label_w - 0.82
        for i, (label, val) in enumerate(items):
            yy = y + i * pitch
            self.text(x, yy + 0.02, label_w - 0.10, 0.30, label, size=SZ_BODY,
                      color=INK)
            self.shape(bx, yy + 0.05, bw, bar_h, fill=TRACK)
            self.shape(bx, yy + 0.05, bw * val / vmax, bar_h,
                       fill=CORAL if i < highlight else BLUE_M)
            self.text(bx + bw + 0.10, yy + 0.03, 0.72, 0.28,
                      "{0}{1}".format(val, unit), size=SZ_BODY, bold=True,
                      color=NAVY)

    def chevron(self, x, y, w=0.32, h=2.10, color=CHEVRON):
        return self.shape(x, y, w, h, MSO_SHAPE.CHEVRON, fill=color)

    def notes(self, body):
        self.sld.notes_slide.notes_text_frame.text = body


# ===================================================================== Deck
class Deck(object):
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.page = 0

    # ------------------------------------------------------------ 内部
    def _blank(self, dark=False):
        return Slide(self, self.prs.slides.add_slide(self.prs.slide_layouts[6]),
                     dark=dark)

    def _gradient(self, s):
        sp = s.sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                    Inches(SW), Inches(SH))
        sp.line.fill.background()
        sp.shadow.inherit = False
        fill = sp.fill
        fill.gradient()
        fill.gradient_angle = 25.0
        fill.gradient_stops[0].color.rgb = NAVY_DP
        fill.gradient_stops[0].position = 0.0
        fill.gradient_stops[1].color.rgb = CORAL
        fill.gradient_stops[1].position = 1.0
        return sp

    def _header(self, s, title, dark=False):
        s.text(ML, HDR_Y, CW - 2.4, 0.42, title, size=SZ_TITLE, bold=True,
               color=WHITE if dark else NAVY)
        if dark:
            s.picture(LOGO_W, 10.82, HDR_Y, h=0.42)
        else:
            s.picture(LOGO, MR - 1.24, HDR_Y, h=0.26)
            s.rule(ML, RULE_Y, CW)

    def _footer(self, s):
        self.page += 1
        s.rule(ML, FOOT_RULE_Y, CW, h=0.012)
        s.text(ML, 7.13, CW - 0.35, 0.20, FOOTER, size=SZ_FOOT, color=GRAY_L,
               align=PP_ALIGN.RIGHT)
        s.text(MR - 0.32, 7.11, 0.32, 0.22, str(self.page), size=9, bold=True,
               color=NAVY, align=PP_ALIGN.RIGHT)

    # ------------------------------------------------------------ ページ種別
    def slide(self, title, lead=None):
        """白地の本文ページ（ヘッダー＋リード＋フッター）"""
        s = self._blank()
        self._header(s, title)
        self._footer(s)
        if lead:
            s.lead(lead)
        return s

    def gradient_slide(self, title, lead=None):
        """グラデーション面のページ（締めやセクション扉に使う。フッター無し）"""
        s = self._blank(dark=True)
        self._gradient(s)
        self._header(s, title, dark=True)
        self.page += 1
        if lead:
            s.lead(lead)
        return s

    def cover(self, headline, subline, eyebrow=None, supports=(), footnote=None,
              right_note=None):
        """表紙"""
        s = self._blank(dark=True)
        self._gradient(s)
        s.picture(LOGO_W, ML, 0.62, h=0.42)
        self.page += 1
        if eyebrow:
            s.text(ML, 1.85, 9.6, 0.34, eyebrow, size=13, bold=True,
                   color=ROSE_L)
        s.text(ML, 2.42, 10.6, 2.1,
               [(headline, {"size": 40, "bold": True}),
                (subline, {"size": 26, "bold": True})],
               color=WHITE, spacing=1.22)
        s.shape(ML, 4.72, 1.5, 0.05, fill=WHITE)
        if supports:
            s.text(ML, 5.02, 10.2, 0.9,
                   [(t, {"size": 15 if i == 0 else 13, "bold": i == 0})
                    for i, t in enumerate(supports)],
                   color=CREAM, spacing=1.5)
        if footnote:
            s.text(ML, 6.62, 6.0, 0.3, footnote, size=11, color=CREAM)
        if right_note:
            s.text(7.6, 6.62, 5.1, 0.3, right_note, size=10, color=CREAM,
                   align=PP_ALIGN.RIGHT)
        return s

    def contact(self, headline_lines, company, contacts, eyebrow=None):
        """裏表紙・お問い合わせ（今回のデッキでは未使用。必要なら1行で足せる）"""
        s = self._blank(dark=True)
        self._gradient(s)
        s.picture(LOGO_W, ML, 0.68, h=0.42)
        self.page += 1
        if eyebrow:
            s.text(ML, 2.30, 9.0, 0.34, eyebrow, size=13, bold=True,
                   color=ROSE_L)
        s.text(ML, 2.86, 10.4, 1.5,
               [(t, {"size": 34, "bold": True}) for t in headline_lines],
               color=WHITE, spacing=1.2)
        s.shape(ML, 4.88, 1.5, 0.05, fill=WHITE)
        s.text(ML, 5.22, 6.2, 0.9,
               [(company, {"size": 14, "bold": True}),
                (contacts, {"size": 12})], color=CREAM, spacing=1.5)
        return s

    def save(self, path):
        self.prs.save(path)
        return path
